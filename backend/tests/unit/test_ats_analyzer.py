import pytest
import pytest_asyncio
import uuid
from httpx import AsyncClient, ASGITransport
from unittest.mock import patch, MagicMock
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker, AsyncSession

from app.main import app
from app.core.database import Base, get_db
from app.models.user import User, UserRole
from app.models.resume import Resume
from app.models.job import Job
from app.api.dependencies.auth import get_current_active_user
from app.services.ats_analyzer import analyze_ats, _gemini_analysis

TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"

student_1 = User(
    id=uuid.UUID("22222222-2222-2222-2222-222222222222"),
    email="student1@university.edu",
    hashed_password="hash",
    full_name="Alex Rivera",
    role=UserRole.STUDENT,
    is_active=True,
    is_verified=True,
)

student_2 = User(
    id=uuid.UUID("33333333-3333-3333-3333-333333333333"),
    email="student2@university.edu",
    hashed_password="hash",
    full_name="Sam Taylor",
    role=UserRole.STUDENT,
    is_active=True,
    is_verified=True,
)


@pytest_asyncio.fixture(scope="function")
async def db_session():
    engine = create_async_engine(TEST_DATABASE_URL, echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    session_factory = async_sessionmaker(
        bind=engine,
        class_=AsyncSession,
        expire_on_commit=False,
    )

    async with session_factory() as session:
        yield session

    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.drop_all)
    await engine.dispose()


@pytest_asyncio.fixture(scope="function")
async def client(db_session):
    async def override_get_db():
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.mark.asyncio
async def test_unauthenticated_analyze_rejected(client):
    """Unauthenticated request to analyze-ats returns 401."""
    resp = await client.post("/api/resumes/analyze-ats", json={"resume_text": "...", "job_description": "..."})
    assert resp.status_code == 401


@pytest.mark.asyncio
async def test_missing_inputs_rejected(client, db_session):
    """Returns 400 if resume or JD inputs are missing."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1

    # Missing job description
    resp1 = await client.post("/api/resumes/analyze-ats", json={"resume_text": "Experienced Python Engineer"})
    assert resp1.status_code == 400

    # Missing resume text
    resp2 = await client.post("/api/resumes/analyze-ats", json={"job_description": "Seeking Python Engineer"})
    assert resp2.status_code == 400


@pytest.mark.asyncio
async def test_analyze_with_raw_text(client):
    """Student can analyze raw resume text against custom job description."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1

    payload = {
        "resume_text": (
            "Alex Rivera\n"
            "Technical Skills: Python, FastAPI, React, Docker, PostgreSQL\n"
            "Experience: Software Engineer at NovaTech. Architected APIs cutting latency by 35%.\n"
            "Education: B.S. in Computer Science\n"
        ),
        "job_description": (
            "Job Title: Backend Developer\n"
            "Requirements: Strong experience in Python, FastAPI, PostgreSQL, Kubernetes, and AWS.\n"
        ),
    }

    resp = await client.post("/api/resumes/analyze-ats", json=payload)
    assert resp.status_code == 200
    data = resp.json()

    assert "overall_score" in data
    assert 0 <= data["overall_score"] <= 100
    assert "verdict" in data
    assert "breakdown" in data
    assert data["breakdown"]["skills_score"] > 0
    assert "matched_skills" in data
    assert "missing_skills" in data
    # Python & FastAPI & PostgreSQL should be matched
    matched_lower = [m.lower() for m in data["matched_skills"]]
    assert "python" in matched_lower
    assert "fastapi" in matched_lower
    assert "postgresql" in matched_lower
    # Kubernetes & AWS should be missing
    missing_lower = [m.lower() for m in data["missing_skills"]]
    assert "kubernetes" in missing_lower
    assert "aws" in missing_lower
    assert len(data["bullet_recommendations"]) >= 1


@pytest.mark.asyncio
async def test_analyze_with_saved_resume_and_job_id(client, db_session):
    """Student can analyze their saved resume entity against an active platform job listing."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1
    db_session.add(student_1)

    # 1. Create a saved resume for student 1
    resume = Resume(
        user_id=student_1.id,
        title="Full Stack Resume",
        template_id="classic-tech",
        personal_info={"full_name": "Alex Rivera", "email": "alex@example.com"},
        summary="Experienced Full Stack Engineer with focus on Python and React.",
        skills=["Python", "FastAPI", "React", "PostgreSQL", "Docker"],
        experience=[
            {
                "role": "Frontend Intern",
                "company": "Acme Corp",
                "bullets": ["Optimized page rendering by 40%."],
            }
        ],
        education=[{"institution": "UW", "degree": "B.S. in CS"}],
        projects=[{"name": "App", "tech_stack": "Python, React"}],
        is_primary=True,
    )
    db_session.add(resume)

    # 2. Create a platform job
    job = Job(
        title="Cloud Software Engineer",
        company_name="CloudNova Systems",
        domain="Software Engineering",
        job_type="Full-time",
        work_mode="Remote",
        location="Remote",
        apply_url="https://cloudnova.com/jobs/1",
        skills_required=["Python", "FastAPI", "Docker", "AWS", "Kubernetes"],
        description="Looking for an engineer to build distributed systems in Python and FastAPI.",
        responsibilities=["Develop microservices", "Manage cloud deployments"],
        requirements=["3+ years experience with Python", "Hands on Docker and AWS"],
        is_active=True,
    )
    db_session.add(job)
    await db_session.commit()
    await db_session.refresh(resume)
    await db_session.refresh(job)

    # 3. Analyze using UUIDs
    resp = await client.post(
        "/api/resumes/analyze-ats",
        json={"resume_id": str(resume.id), "job_id": str(job.id)},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["overall_score"] >= 40
    matched_lower = [m.lower() for m in data["matched_skills"]]
    assert "python" in matched_lower
    assert "fastapi" in matched_lower
    assert "docker" in matched_lower


@pytest.mark.asyncio
async def test_cross_user_resume_analysis_isolated(client, db_session):
    """Student 2 cannot run analysis against Student 1's saved resume."""
    db_session.add_all([student_1, student_2])
    resume = Resume(
        user_id=student_1.id,
        title="Alex Rivera Private Resume",
        template_id="modern-minimal",
        personal_info={"full_name": "Alex Rivera"},
    )
    db_session.add(resume)
    await db_session.commit()
    await db_session.refresh(resume)

    # Act as Student 2
    app.dependency_overrides[get_current_active_user] = lambda: student_2
    resp = await client.post(
        "/api/resumes/analyze-ats",
        json={"resume_id": str(resume.id), "job_description": "General Software Engineer"},
    )
    assert resp.status_code == 404


@pytest.mark.asyncio
async def test_extract_resume_text_valid_docx(client):
    """Student can upload a valid .docx file and extract clean resume text."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1
    import io
    import docx

    # Generate an in-memory DOCX
    doc = docx.Document()
    doc.add_paragraph("Alex Rivera - Full Stack Developer")
    doc.add_paragraph("Skills: Python, FastAPI, React, Docker, PostgreSQL")
    doc.add_paragraph("Experience: Built scalable cloud APIs and deployed with Docker.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "University of California"
    table.rows[0].cells[1].text = "B.S. Computer Science"

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    files = {"file": ("my_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    resp = await client.post("/api/resumes/extract-text", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["filename"] == "my_resume.docx"
    assert "Alex Rivera" in data["extracted_text"]
    assert "FastAPI" in data["extracted_text"]
    assert "B.S. Computer Science" in data["extracted_text"]
    assert data["character_count"] > 50
    assert data["word_count"] > 10
    assert len(data["preview"]) > 0


@pytest.mark.asyncio
async def test_extract_resume_text_valid_pdf(client):
    """Student can upload a valid .pdf file and extract clean resume text."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1
    import io
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    c.drawString(100, 750, "Jordan Lee - Backend Systems Engineer")
    c.drawString(100, 730, "Core Skills: Python, FastAPI, Redis, Kubernetes, AWS")
    c.drawString(100, 710, "Built microservices handling 10k requests per second.")
    c.save()
    pdf_bytes = buffer.getvalue()

    files = {"file": ("jordan_resume.pdf", pdf_bytes, "application/pdf")}
    resp = await client.post("/api/resumes/extract-text", files=files)
    assert resp.status_code == 200
    data = resp.json()
    assert data["filename"] == "jordan_resume.pdf"
    assert "Jordan Lee" in data["extracted_text"]
    assert "FastAPI" in data["extracted_text"]


@pytest.mark.asyncio
async def test_extract_resume_unsupported_format(client):
    """Rejects files with unsupported extensions or MIME types."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1
    files = {"file": ("notes.txt", b"Plain text resume content", "text/plain")}
    resp = await client.post("/api/resumes/extract-text", files=files)
    assert resp.status_code == 400
    assert "Unsupported file format" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_extract_resume_oversized_file(client):
    """Rejects files exceeding 5 MB limit."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1
    huge_bytes = b"%PDF-1.4 " + b"0" * (6 * 1024 * 1024)
    files = {"file": ("huge_resume.pdf", huge_bytes, "application/pdf")}
    resp = await client.post("/api/resumes/extract-text", files=files)
    assert resp.status_code == 400
    assert "exceeds maximum allowed limit" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_extract_resume_corrupted_pdf(client):
    """Gracefully handles corrupted PDF documents."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1
    corrupt_bytes = b"%PDF-1.4 completely corrupted non-renderable bytes\x00\x01\x02"
    files = {"file": ("corrupted.pdf", corrupt_bytes, "application/pdf")}
    resp = await client.post("/api/resumes/extract-text", files=files)
    assert resp.status_code == 400
    assert "Corrupted or invalid PDF" in resp.json()["detail"] or "unable to parse" in resp.json()["detail"].lower()


@pytest.mark.asyncio
async def test_extract_resume_empty_scanned_pdf(client):
    """Gracefully handles PDFs without extractable text (e.g. image-only / blank)."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1
    import io
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    # Don't draw any text strings — empty blank page
    c.showPage()
    c.save()
    pdf_bytes = buffer.getvalue()

    files = {"file": ("scanned_blank.pdf", pdf_bytes, "application/pdf")}
    resp = await client.post("/api/resumes/extract-text", files=files)
    assert resp.status_code == 400
    assert "readable text" in resp.json()["detail"].lower()


@pytest.mark.asyncio
async def test_end_to_end_extracted_resume_to_ats_analysis(client):
    """Verifies that text extracted from an uploaded document feeds directly into ATS analysis."""
    app.dependency_overrides[get_current_active_user] = lambda: student_1
    import io
    import docx

    # 1. Upload DOCX resume
    doc = docx.Document()
    doc.add_paragraph("Morgan Developer")
    doc.add_paragraph("Skills: Python, FastAPI, Docker, PostgreSQL, Redis, Unit Testing")
    doc.add_paragraph("Experience: Developed backend microservices using FastAPI and Docker.")
    buffer = io.BytesIO()
    doc.save(buffer)

    files = {"file": ("morgan_cv.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    extract_resp = await client.post("/api/resumes/extract-text", files=files)
    assert extract_resp.status_code == 200
    extracted_text = extract_resp.json()["extracted_text"]

    # 2. Submit extracted text to ATS analysis pipeline
    analysis_resp = await client.post(
        "/api/resumes/analyze-ats",
        json={
            "resume_text": extracted_text,
            "job_description": "Seeking Python Engineer with FastAPI and Docker experience.",
        },
    )
    assert analysis_resp.status_code == 200
    ats_data = analysis_resp.json()
    assert ats_data["overall_score"] >= 40
    matched = [m.lower() for m in ats_data["matched_skills"]]
    assert "python" in matched
    assert "fastapi" in matched
    assert "docker" in matched

